# -*- coding: utf-8 -*-
"""组会文献汇报 PPT: 对抗性场景生成文献 + GA2AD 框架解读"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "组会文献汇报_GA2AD_20260901.pptx")

DARK_BLUE   = RGBColor(0x1B, 0x3A, 0x5C)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
ACCENT_RED  = RGBColor(0xE7, 0x4C, 0x3C)
ACCENT_GREEN= RGBColor(0x27, 0xAE, 0x60)
ORANGE      = RGBColor(0xE6, 0x7E, 0x22)
DARK_GRAY   = RGBColor(0x2C, 0x3E, 0x50)
LIGHT_GRAY  = RGBColor(0x95, 0xA5, 0xA6)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)


def title_bar(slide, title, subtitle=None):
    s = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.9))
    s.fill.solid(); s.fill.fore_color.rgb = DARK_BLUE; s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(25); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.6); tf.margin_top = Inches(0.13)
    if subtitle:
        p2 = tf.add_paragraph(); p2.text = subtitle; p2.font.size = Pt(12); p2.font.color.rgb = RGBColor(0xAE,0xD6,0xF1); p2.alignment = PP_ALIGN.LEFT


def bullets(slide, left, top, w, h, items, fs=Pt(15)):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b; p.font.size = fs; p.font.color.rgb = DARK_GRAY; p.space_after = Pt(6)
    return tb


def highlight_box(slide, left, top, w, h, text, color=ACCENT_RED, fs=Pt(14), bg=RGBColor(0xFD,0xED,0xEC)):
    s = slide.shapes.add_shape(1, left, top, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = bg; s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.size = fs; p.font.color.rgb = color; p.font.bold = True; p.alignment = PP_ALIGN.CENTER


def build():
    prs = Presentation()
    prs.slide_width = Inches(10); prs.slide_height = Inches(5.63)
    blank = prs.slide_layouts[6]

    # ===== S1: 标题 =====
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(5.63))
    bg.fill.solid(); bg.fill.fore_color.rgb = DARK_BLUE; bg.line.fill.background()
    tb = s.shapes.add_textbox(Inches(1), Inches(1.1), Inches(8), Inches(1.2)); tf = tb.text_frame
    p = tf.paragraphs[0]; p.text = "文献调研：对抗性场景生成与 GA2AD 框架"; p.font.size = Pt(30); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
    tb2 = s.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(7), Inches(0.8)); tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]; p2.text = "AEB / 自动驾驶安全测试三大方法 + 导师推荐论文 GA2AD 解读"; p2.font.size = Pt(16); p2.font.color.rgb = RGBColor(0xAE,0xD6,0xF1); p2.alignment = PP_ALIGN.CENTER
    tb3 = s.shapes.add_textbox(Inches(2), Inches(3.5), Inches(6), Inches(0.6)); tf3 = tb3.text_frame
    p3 = tf3.paragraphs[0]; p3.text = "导师组会汇报  |  2026年9月1日"; p3.font.size = Pt(14); p3.font.color.rgb = LIGHT_GRAY; p3.alignment = PP_ALIGN.CENTER

    # ===== S2: 文献全景 =====
    s = prs.slides.add_slide(blank)
    title_bar(s, "文献全景：安全测试的三大方法", "彭老师要求检索 AEB testing / risk assessment，现存方法可归为三类")
    items = [
        "① 形式化验证（我目前的方法）",
        "   Petri 网 / UPPAAL / SOTIF：确定性穷举，能保证找到碰撞",
        "   代表：AEB 着色混合 Petri 网验证；COSACC 用 UPPAAL 做时序验证",
        "",
        "② 对抗性场景生成（彭老师说的 adversarial ML）",
        "   RL / GAN / 遗传算法：主动学“最危险的前车行为”",
        "   代表：WGAN 生成切入场景；PAFOT 遗传算法搜最优测试",
        "",
        "③ 贝叶斯优化安全评估（彭老师建议的方向）",
        "   高斯过程代理模型 + acquisition function：小样本估碰撞概率",
        "   代表：Hao Wu 2026 用贝叶斯优化做 cut-in 场景碰撞概率估计",
    ]
    bullets(s, Inches(0.7), Inches(1.1), Inches(8.6), Inches(3.4), items, Pt(14))
    highlight_box(s, Inches(0.7), Inches(4.5), Inches(8.6), Inches(0.6),
                  "关键洞察：①有确定性无效率，②③有效率无保证 —— 把“形式化验证当 oracle + 贝叶斯优化选点”是空白点")

    # ===== S3: GA2AD 框架 =====
    s = prs.slides.add_slide(blank)
    title_bar(s, "GA2AD：自适应对抗训练框架（导师推荐）", "Xu et al., Nature Communications 2026 —— 把 AV 当黑盒，训练背景车主动制造危险")
    # 三个模块卡片
    cards = [
        ("① 混合风险场 + 贝叶斯", "传统风险场 + TTC\n→ 危险/注意/安全三级概率\n连续风险值 ϵ 引导 agent", RGBColor(0xD6,0xEA,0xF8), ACCENT_BLUE),
        ("② 激活函数", "决定“何时”施加对抗动作\n卡在 [Ml, Mu] 区间\n避免太弱(练不出)或太猛(练崩)", RGBColor(0xFD,0xED,0xEC), ACCENT_RED),
        ("③ 自适应切换模块", "按碰撞间隔交替训练\n目标车(TV) ↔ 背景车(BV)\n双智能体→单智能体，稳定训练", RGBColor(0xEA,0xFA,0xF1), ACCENT_GREEN),
    ]
    x = 0.4
    for title, body, fill, color in cards:
        box = s.shapes.add_shape(1, Inches(x), Inches(1.2), Inches(3.0), Inches(2.4))
        box.fill.solid(); box.fill.fore_color.rgb = fill; box.line.fill.background()
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = color; p.alignment = PP_ALIGN.CENTER
        for line in body.split("\n"):
            p2 = tf.add_paragraph(); p2.text = line; p2.font.size = Pt(11); p2.font.color.rgb = DARK_GRAY; p2.alignment = PP_ALIGN.CENTER
        x += 3.15
    # 结果条
    highlight_box(s, Inches(0.4), Inches(3.9), Inches(9.2), Inches(0.6),
                  "结果：碰撞率 ↓55–70%，jerk ↓21–42%，黑盒适配 DQN / TD3 / SAC / DDPG", ACCENT_GREEN, Pt(14), RGBColor(0xEA,0xFA,0xF1))
    bullets(s, Inches(0.4), Inches(4.6), Inches(9.2), Inches(0.9),
            ["定位：这是“防御侧”——训练 AV 变得更鲁棒；不是“找危险场景”。"], Pt(13))

    # ===== S4: 对我的启发 =====
    s = prs.slides.add_slide(blank)
    title_bar(s, "GA2AD 对我的启发：从“搜参数”到“学策略”", "我的工作对应 GA2AD 里的“背景车模型(BV)”——进攻侧")
    items = [
        "定位：GA2AD=防御侧(训练AV鲁棒)，我=进攻侧(找危险场景)＝它的 BV 模型",
        "",
        "三个可借鉴：",
        "  ① 混合风险场+贝叶斯 → 替代我的二元奖励(碰撞+10/安全-1)，收敛更快",
        "  ② 激活函数 → 控制对抗强度，避免太保守/太激进",
        "  ③ 学“行为策略” → 我现在搜静态参数(距离/速度)，应升级为前车动作序列",
        "",
        "回扣彭老师纠正：“学的不是 policy，是参数 vector”——GA2AD 展示了学 policy 的样子",
    ]
    bullets(s, Inches(0.7), Inches(1.1), Inches(8.6), Inches(3.4), items, Pt(14))
    highlight_box(s, Inches(0.7), Inches(4.5), Inches(8.6), Inches(0.6),
                  "下一步：前车行为学成策略(何时刹车/切入)，用风险场做奖励，UPPAAL 当 oracle 验证")

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == '__main__':
    print("PPTX saved:", build())
