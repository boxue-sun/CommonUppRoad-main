# -*- coding: utf-8 -*-
"""Generate clean 8-slide PPT for 2026-06-30 group meeting."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

ROOT = os.path.dirname(__file__)
OUT = os.path.join(ROOT, "experiments", "meeting_20260630_clean.pptx")

DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
ACCENT_BLUE = RGBColor(0x2E, 0x86, 0xC1)
ACCENT_RED = RGBColor(0xE7, 0x4C, 0x3C)
ACCENT_GREEN = RGBColor(0x27, 0xAE, 0x60)
DARK_GRAY = RGBColor(0x2C, 0x3E, 0x50)
LIGHT_GRAY = RGBColor(0x95, 0xA5, 0xA6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def sn(slide, n, t=8):
    tb = slide.shapes.add_textbox(Inches(8.5), Inches(5.2), Inches(1.0), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "%d/%d" % (n, t); p.font.size = Pt(10); p.font.color.rgb = LIGHT_GRAY; p.alignment = PP_ALIGN.RIGHT


def tb(slide, t, sub=None):
    s = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.9))
    s.fill.solid(); s.fill.fore_color.rgb = DARK_BLUE; s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = t; p.font.size = Pt(26); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.6); tf.margin_top = Inches(0.15)
    if sub:
        p2 = tf.add_paragraph(); p2.text = sub; p2.font.size = Pt(13); p2.font.color.rgb = RGBColor(0xAE, 0xD6, 0xF1)


def bl(slide, l, t, w, h, items, fs=Pt(16)):
    bx = slide.shapes.add_textbox(l, t, w, h)
    tf = bx.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.size = fs; p.font.color.rgb = DARK_GRAY; p.space_after = Pt(8)
    return bx


def hb(slide, l, t, w, h, text, c=ACCENT_RED):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor(0xFD, 0xED, 0xEC); s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(15); p.font.color.rgb = c; p.font.bold = True; p.alignment = PP_ALIGN.CENTER


def img(slide, path, l, t, w, h=None):
    if os.path.exists(path):
        if h: return slide.shapes.add_picture(path, l, t, w, h)
        else: return slide.shapes.add_picture(path, l, t, width=w)
    return None


def build():
    prs = Presentation()
    prs.slide_width = Inches(10); prs.slide_height = Inches(5.63)
    B = prs.slide_layouts[6]

    # ===== S1: Title =====
    s = prs.slides.add_slide(B)
    bg = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(5.63))
    bg.fill.solid(); bg.fill.fore_color.rgb = DARK_BLUE; bg.line.fill.background()
    bx = s.shapes.add_textbox(Inches(1), Inches(1.2), Inches(8), Inches(1.2))
    tf = bx.text_frame
    p = tf.paragraphs[0]; p.text = "ISO 15622 ACC模型验证 + RL危险场景搜索"; p.font.size = Pt(30); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
    bx2 = s.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(7), Inches(0.8))
    tf2 = bx2.text_frame
    p2 = tf2.paragraphs[0]; p2.text = "从IDM到法规模型 + 实验验证 + 新想法探讨"; p2.font.size = Pt(17); p2.font.color.rgb = RGBColor(0xAE, 0xD6, 0xF1); p2.alignment = PP_ALIGN.CENTER
    bx3 = s.shapes.add_textbox(Inches(2), Inches(3.8), Inches(6), Inches(0.5))
    p3 = bx3.text_frame.paragraphs[0]; p3.text = "导师组会汇报  |  2026年6月30日"; p3.font.size = Pt(14); p3.font.color.rgb = LIGHT_GRAY; p3.alignment = PP_ALIGN.CENTER

    # ===== S2: From last time to this time =====
    s = prs.slides.add_slide(B)
    tb(s, "上次的问题 -> 这次做了两件事", "")
    sn(s, 2)

    bl(s, Inches(0.5), Inches(1.2), Inches(4.5), Inches(3.8), [
        "上次的问题:",
        "  x IDM 模型没有反应延迟，立刻响应",
        "  x 加速度瞬间跳变 (0 -> -19 m/s^2)",
        "  x 碰撞场景不现实",
        "",
        "这次做了两件事:",
        "",
        "第一件事 (已完成):",
        "  换模型 + 跑实验",
        "  IDM -> ISO 15622 四模式ACC",
        "  重新跑 200变体 + 新增急刹实验",
    ], Pt(13))

    bl(s, Inches(5.3), Inches(1.2), Inches(4.5), Inches(3.8), [
        "第二件事 (新想法):",
        "  RL 搜索参数空间",
        "",
        "  问题: 参数多了网格穷举行不通",
        "  想法: 让 RL Agent 学会",
        "  '什么参数容易碰撞'",
        "",
        "  今天重点: 讲第一件事的结果",
        "  然后说第二件事的想法",
        "  请老师帮把关方向",
    ], Pt(13))

    # ===== S3: ISO 15622 model =====
    s = prs.slides.add_slide(B)
    tb(s, "ISO 15622 四模式 ACC", "国际标准跟车模型，加了 1.5 秒反应延迟")
    sn(s, 3)

    # Four mode boxes
    modes = [
        ("自由流", "前面没车\n-> 加速到 30m/s", RGBColor(0xEA, 0xFA, 0xF1)),
        ("接近", "前面有车但远\n-> 慢慢减速", RGBColor(0xFE, 0xF9, 0xE7)),
        ("跟车", "保持 1.8s 车距\n-> 稳速跟随", RGBColor(0xEB, 0xF5, 0xFB)),
        ("制动", "太近了!\n-> 急刹 -6m/s^2", RGBColor(0xFD, 0xED, 0xEC)),
    ]
    for i, (name, desc, clr) in enumerate(modes):
        x = Inches(0.5 + i * 2.4)
        box = s.shapes.add_shape(1, x, Inches(1.3), Inches(2.1), Inches(1.6))
        box.fill.solid(); box.fill.fore_color.rgb = clr; box.line.color.rgb = ACCENT_BLUE; box.line.width = Pt(1.5)
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = name; p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = DARK_BLUE; p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph(); p2.text = desc; p2.font.size = Pt(11); p2.font.color.rgb = DARK_GRAY; p2.alignment = PP_ALIGN.CENTER

    bl(s, Inches(0.8), Inches(3.3), Inches(8.5), Inches(1.8), [
        "和 IDM 的最大区别: 反应延迟 1.5 秒",
        "  IDM = 注意力超级集中的老司机，前车一有动作立刻反应",
        "  ISO 15622 = 开了 ACC 的普通司机，检测到危险后 1.5 秒才开始刹车",
        "  1.5 秒在高速 (30m/s) 下 = 跑出去 45 米才开始刹！",
    ], Pt(14))

    # ===== S4: Experiment 1 - heatmap =====
    s = prs.slides.add_slide(B)
    tb(s, "实验1: 200变体碰撞热力图", "8距离 x 5前车速度 x 5ego速度 = 200次UPPAAL验证")
    sn(s, 4)

    img(s, os.path.join(ROOT, "experiments", "heatmap_analysis.png"), Inches(0.3), Inches(1.0), Inches(6.2), Inches(3.5))

    bl(s, Inches(6.8), Inches(1.1), Inches(3.0), Inches(3.8), [
        "结果: 68/200 碰撞",
        "碰撞率: 34%",
        "",
        "安全边界在 15~20m",
        "",
        "三分区:",
        "<=5m  必撞(100%)",
        "8-15m 高速才撞",
        ">=20m 完全安全(0%)",
    ], Pt(13))

    # ===== S5: Experiment 2 - sudden brake =====
    s = prs.slides.add_slide(B)
    tb(s, "实验2: 前车急刹场景", "ego和obs同速30m/s -> obs在3s急刹 -> ego 1.5s后反应")
    sn(s, 5)

    bl(s, Inches(0.8), Inches(1.3), Inches(5.0), Inches(3.5), [
        "场景: ego 和 obs 都以 30m/s 行驶",
        "       obs 突然急刹 -10 m/s^2",
        "       ego 有 1.5s 反应延迟",
        "",
        "结果: 22变体中 18个碰撞 (81.8%)",
        "      即使舒适刹车(-4m/s^2)也撞",
        "",
        "原因: 反应延迟 1.5 秒内",
        "      ego 还在全速前进",
        "      前车已经在减速了",
    ], Pt(14))

    bl(s, Inches(6.0), Inches(1.3), Inches(3.8), Inches(3.5), [
        "两个实验总结:",
        "",
        "1. 正常跟车:",
        "   20m以上绝对安全",
        "",
        "2. 前车急刹:",
        "   几乎是必撞场景",
        "   1.5s反应延迟是主因",
    ], Pt(14))

    # ===== S6: Problem -> RL motivation =====
    s = prs.slides.add_slide(B)
    tb(s, "新问题: 参数多了怎么办?", "从网格穷举 -> 智能搜索")
    sn(s, 6)

    bl(s, Inches(0.8), Inches(1.3), Inches(8.5), Inches(1.8), [
        "目前实验只有 3 个参数 (距离, ego速度, obs速度) -> 200 种组合，能穷举",
        "但如果加入: 反应时间, 刹车强度, 切入角度, 路面摩擦...",
        "  3维 x 8值 = 512 种  -> 能跑",
        "  4维 x 8值 = 4,096  -> 勉强",
        "  5维 x 8值 = 32,768 -> 跑不完",
        "  6维 x 8值 = 262,144 -> 不可行",
    ], Pt(14))

    hb(s, Inches(0.8), Inches(3.5), Inches(8.5), Inches(0.8),
        "想法: 不要让机器穷举，让机器学会 '什么参数容易碰撞'，优先搜那些区域")

    bl(s, Inches(0.8), Inches(4.5), Inches(8.5), Inches(0.8), [
        "-> 这就是强化学习要做的事: 把参数搜索变成一个学习问题",
    ], Pt(14))

    # ===== S7: RL Design =====
    s = prs.slides.add_slide(B)
    tb(s, "RL 设计: Agent 在参数空间中学会 '哪里危险'", "三个关键元素: 状态 / 动作 / 奖励")
    sn(s, 7)

    # Three boxes
    items = [
        ("状态 (State)", "Agent 现在在哪?", "当前参数组合\n(距离, ego速度, obs速度)\n= 3D棋盘的坐标", RGBColor(0xD6, 0xEA, 0xF8)),
        ("动作 (Action)", "Agent 往哪走?", "6种:\n距离+1, 距离-1,\nego+1, ego-1,\nobs+1, obs-1", RGBColor(0xEA, 0xFA, 0xF1)),
        ("奖励 (Reward)", "走对了还是错了?", "碰撞 -> +10\n安全 -> -1\n重复 -> -0.5 (别打转)", RGBColor(0xFD, 0xED, 0xEC)),
    ]
    for i, (title, q, detail, clr) in enumerate(items):
        x = Inches(0.5 + i * 3.2)
        box = s.shapes.add_shape(1, x, Inches(1.3), Inches(2.9), Inches(2.2))
        box.fill.solid(); box.fill.fore_color.rgb = clr; box.line.fill.background()
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = DARK_BLUE; p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph(); p2.text = q; p2.font.size = Pt(11); p2.font.color.rgb = ACCENT_BLUE; p2.alignment = PP_ALIGN.CENTER
        p3 = tf.add_paragraph(); p3.text = "\n" + detail; p3.font.size = Pt(12); p3.font.color.rgb = DARK_GRAY; p3.alignment = PP_ALIGN.CENTER

    # Training curve
    tc = os.path.join(ROOT, "experiments", "rl_search", "training_curve.png")
    img(s, tc, Inches(0.5), Inches(3.8), Inches(5.5), Inches(1.7))
    bl(s, Inches(6.3), Inches(3.8), Inches(3.5), Inches(1.7), [
        "训练曲线:",
        "上: 奖励一直在涨 ->",
        "    Agent 确实在学",
        "下: 碰撞数在增加 ->",
        "    效率越来越高",
    ], Pt(12))

    # ===== S8: Ask for feedback =====
    s = prs.slides.add_slide(B)
    bg = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(5.63))
    bg.fill.solid(); bg.fill.fore_color.rgb = DARK_BLUE; bg.line.fill.background()
    sn(s, 8)

    bx = s.shapes.add_textbox(Inches(1), Inches(0.6), Inches(8), Inches(0.7))
    p = bx.text_frame.paragraphs[0]; p.text = "请老师帮我看方向"; p.font.size = Pt(28); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER

    bl(s, Inches(1.5), Inches(1.6), Inches(7), Inches(3.5), [
        "目前状态:",
        "  RL是离线的 -> Agent从已有200变体数据里学",
        "  验证了 'RL能学会碰撞规律' (训练曲线在涨)",
        "",
        "我想做的:",
        "  改成在线 -> Agent直接调UPPAAL",
        "  在更大的、未知的参数空间里搜索",
        "  参数越多，优势越明显",
        "",
        "想请老师帮我看:",
        "  1. 这个方向对不对? 值不值得做?",
        "  2. 继续Q-learning还是换DQN/贝叶斯优化?",
        "  3. 有推荐的RL算法或论文吗?",
    ], Pt(14))

    prs.save(OUT)
    return OUT


if __name__ == '__main__':
    p = build()
    print("Saved:", p)
